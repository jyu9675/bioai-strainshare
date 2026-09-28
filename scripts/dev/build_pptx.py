#!/usr/bin/env python3
"""Build the strainshare progress-update deck (mentor presentation)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

ACCENT = RGBColor(0x9C, 0x2B, 0x39)   # clinical garnet
TEAL   = RGBColor(0x2C, 0x5A, 0x5F)   # culture teal
INK    = RGBColor(0x24, 0x1A, 0x20)
MUTED  = RGBColor(0x6E, 0x61, 0x67)
GROUND = RGBColor(0xF8, 0xF5, 0xF3)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
FIG = "/mnt/c/Jeanyu/BIOAI/example/cervix_vagina_n26/figures/fig1_within_vs_between.png"
OUT = "/mnt/c/Jeanyu/BIOAI/docs/strainshare_progress.pptx"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height

def bg(slide, color=GROUND):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color

def bar(slide):  # left accent bar
    b = slide.shapes.add_shape(1, 0, 0, Inches(0.18), SH)
    b.fill.solid(); b.fill.fore_color.rgb = ACCENT; b.line.fill.background()

def tb(slide, x, y, w, h):
    box = slide.shapes.add_textbox(x, y, w, h); box.text_frame.word_wrap = True
    return box.text_frame

def para(tf, text, size, color, bold=False, first=False, space=6, level=0, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.level = level; p.space_after = Pt(space)
    r = p.add_run(); r.text = text
    f = r.font; f.size = Pt(size); f.bold = bold; f.italic = italic
    f.color.rgb = color; f.name = "Calibri"
    return p

def title(slide, text, eyebrow=None):
    if eyebrow:
        t = tb(slide, Inches(0.6), Inches(0.35), Inches(11.5), Inches(0.4))
        para(t, eyebrow.upper(), 12, ACCENT, bold=True, first=True)
    t = tb(slide, Inches(0.6), Inches(0.7), Inches(12.2), Inches(1.0))
    para(t, text, 32, INK, bold=True, first=True)

def bullets(slide, items, x=Inches(0.65), y=Inches(1.9), w=Inches(12.0), h=Inches(5.0), size=18):
    tf = tb(slide, x, y, w, h)
    for i, it in enumerate(items):
        if isinstance(it, tuple):
            txt, lvl = it
        else:
            txt, lvl = it, 0
        p = para(tf, ("• " if lvl == 0 else "– ") + txt, size if lvl == 0 else size-2,
                 INK if lvl == 0 else MUTED, first=(i == 0), space=10, level=lvl)
    return tf

# ---- Slide 1: title ----
s = prs.slides.add_slide(BLANK); bg(s, INK)
t = tb(s, Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.4))
para(t, "strainshare", 60, WHITE, bold=True, first=True)
t = tb(s, Inches(0.95), Inches(3.5), Inches(11.0), Inches(1.2))
para(t, "Contamination-aware detection of shared bacterial strains across body sites", 22,
     RGBColor(0xD9,0xC7,0xCB), first=True)
t = tb(s, Inches(0.95), Inches(5.6), Inches(11.5), Inches(1.4))
para(t, "Progress update  ·  Kwon Lab", 16, RGBColor(0xDE,0x76,0x81), bold=True, first=True)
para(t, "Tool: github.com/jyu9675/bioai-strainshare  ·  DOI 10.5281/zenodo.22275588", 13,
     RGBColor(0xB6,0xA9,0xAF))
para(t, "A validated tool, an honest data investigation, and a fundable study", 13,
     RGBColor(0xB6,0xA9,0xAF), italic=True)

# ---- Slide 2: the question ----
s = prs.slides.add_slide(BLANK); bg(s); bar(s)
title(s, "The question — and the gap", "Motivation")
bullets(s, [
 "Does gut/rectal bacteria seed the vagina as the SAME strain during pregnancy — and does limited water/hygiene/healthcare raise the risk?",
 "Motivated by a lived clinical case; grounded in established GBS biology (the recto-vaginal swab exists for exactly this reason).",
 "Same-strain transfer is already documented for the pathogens — by culture:",
 ("GBS: identical genotype in rectum AND vagina in 18/19 colonized women", 1),
 ("E. coli: ~85% of matched vaginal/urinary isolates are one strain; rectal carriage predicts vaginal", 1),
 ("Gut re-seeds the vagina after intrapartum antibiotics", 1),
 "THE GAP: no one has resolved the same strain moving rectum→vagina in situ, from metagenomes, longitudinally, in a low-resource pregnancy cohort.",
])

# ---- Slide 3: the tool ----
s = prs.slides.add_slide(BLANK); bg(s); bar(s)
title(s, "strainshare — what it does", "The tool")
bullets(s, [
 "Turns inStrain output into standardized, contamination-aware shared-strain calls (popANI ≥ 0.999, breadth ≥ 0.5).",
 "Within- vs between-person NULL — the real signal is a strain shared within a person but not between unrelated people.",
 "Translocation-vs-contamination discriminator (community similarity) — flags cross-sample contamination that mimics sharing.",
 "Generalist filter + longitudinal direction inference (rectum-before-vagina), with 'direction_unresolved' when timing is ambiguous.",
 "Standardized threshold spec — a 'VALENCIA-for-strains' so calls are comparable across studies.",
 "Installable, unit-tested (CI, Py 3.9–3.12), MIT-licensed, public, and DOI-archived.",
])

# ---- Slide 4: validation (with figure) ----
s = prs.slides.add_slide(BLANK); bg(s); bar(s)
title(s, "Validation on public data — it works", "Result")
bullets(s, [
 "Public cohort (ENA PRJNA982400); scaled 8 → 20 → 26 women.",
 "2,137 comparable pairs → 76 within-woman cervix↔vagina shared strains.",
 "Clean null across ~14 species:",
 ("within-person ~90–100% shared", 1),
 ("between-person ~0–6%", 1),
 ("e.g. L. iners 14/14 vs 13/238; G. piotii 8/8 vs 1/113", 1),
 "→ The tool reliably detects TRUE within-person sharing and rejects background.",
], x=Inches(0.65), y=Inches(1.8), w=Inches(6.3), size=17)
if os.path.exists(FIG):
    s.shapes.add_picture(FIG, Inches(7.2), Inches(1.9), width=Inches(5.6))
cap = tb(s, Inches(7.2), Inches(6.6), Inches(5.6), Inches(0.5))
para(cap, "Within- vs between-person popANI, n=26 (public data)", 11, MUTED, italic=True, first=True)

# ---- Slide 5: investigation ----
s = prs.slides.add_slide(BLANK); bg(s); bar(s)
title(s, "Rectovaginal investigation — honest findings", "Biology")
bullets(s, [
 "Ran real data: Fijian rectal+vaginal (n=3) and HMP gut+vaginal (partial, ~7 women).",
 "No within-woman gut↔vaginal shared strain in these small, healthy cohorts — reported correctly as NOT-EVALUABLE (the pathogens weren't vaginally colonized in these women), not as 'no transmission'.",
 "One real positive from longitudinal data: Goltsman subject T18 carries gut L. iners as the SAME strain as her vagina across 10 timepoints.",
 "Literature converges: transmission is real for the pathogens; a hexavalent maternal GBS vaccine (GBS6) is now in Phase 3.",
 "Integrity throughout: never fabricated a positive; every result reported as the data showed it.",
])

# ---- Slide 6: proposal ----
s = prs.slides.add_slide(BLANK); bg(s); bar(s)
title(s, "A fundable study — ready to submit", "Proposal")
bullets(s, [
 "Aim 1 does it happen · Aim 2 which direction/organisms · Aim 3 does it cause disease · Aim 4 who is at risk (WASH).",
 "Design: ~350 pregnant women, low-resource, paired rectal+vaginal deep shotgun, 3–4 timepoints; strainshare as the engine.",
 "Power: WASH arm is binding — ~180 for a 20-pt difference; ~350 enrolled yields ~180–220 evaluable.",
 "Full grant package on GitHub: Specific Aims, Research Strategy, budget/timeline, dbGaP scoping, DMSP, cover letter, letters of support, public brief.",
 "Clone-aware: types shared GBS/E. coli to CC17 / ST131 — directly informs GBS6 vaccine serotype coverage.",
])

# ---- Slide 7: status & asks ----
s = prs.slides.add_slide(BLANK); bg(s); bar(s)
title(s, "Status & asks", "Next")
tf = tb(s, Inches(0.65), Inches(1.8), Inches(12.0), Inches(2.2))
para(tf, "DONE", 15, TEAL, bold=True, first=True, space=4)
for it in ["Tool built, tested, DOI-archived, and validated on 26 women of public data",
           "Honest data investigation + full grant-ready proposal package"]:
    para(tf, "• " + it, 17, INK, space=8)
tf2 = tb(s, Inches(0.65), Inches(4.1), Inches(12.0), Inches(3.0))
para(tf2, "ASKS", 15, ACCENT, bold=True, first=True, space=4)
for it in ["Run strainshare on the lab's 382 paired vaginal–rectal set — the tool is validated; this is the dataset that could test PATHOGEN gut↔vaginal transmission in colonized women",
           "HMP gut↔vaginal at scale on EC2 (~$15/afternoon) — the laptop can't process the heavy stool reliably; runbook is ready",
           "Publish the methods/tool paper — n=26 validation backs it; guidance on scope, venue, authorship"]:
    para(tf2, "• " + it, 17, INK, space=8)

prs.save(OUT)
print("wrote", OUT, "-", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
